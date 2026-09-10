# -*- coding: utf-8 -*-
"""관세청 화장품 수출 데이터 — 실리콘투 워크북 Exports_RawData 시트에서 추출.

Usage:
    python customs_exports_sync.py [워크북경로]   # 생략 시 실리콘투 최신 파일

산출물: data/cosmetics/customs_exports.csv
    year,region,country_en,country_kr,hs,usd_k

- 원출처: 관세청 무역통계 (사용자가 워크북에 정리, 연×국가×HS4, US$000)
- 2026년은 1~5월 YTD 잠정 — 화면에 반드시 표기
- 자동 월간 갱신은 data.go.kr OpenAPI(15100475) 키 발급 후 kbeauty 파이프라인으로
  전환 예정. 그전까지는 워크북이 갱신되면 이 스크립트를 다시 돌린다.
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

WB_DIR = Path(r"C:\Users\user99i1\LK자산운용\LK자산운용 - 문서\Companies\실리콘투")
OUT = Path(__file__).resolve().parent / "data" / "cosmetics" / "customs_exports.csv"


def latest_workbook() -> Path:
    cands = sorted(WB_DIR.glob("실리콘투_Analysis template_*.xlsx"),
                   key=lambda p: p.stat().st_mtime)
    if not cands:
        raise SystemExit("실리콘투 워크북을 찾지 못했습니다")
    return cands[-1]


def open_workbook(path: Path):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        import shutil
        import tempfile
        tmp = Path(tempfile.gettempdir()) / "s2_customs_copy.xlsx"
        try:
            shutil.copy(path, tmp)
        except PermissionError:
            import win32com.client as win32
            xl = win32.GetActiveObject("Excel.Application")
            for w in xl.Workbooks:
                if w.Name == path.name:
                    w.SaveCopyAs(str(tmp))
                    break
            else:
                raise
        print(f"워크북이 잠겨 있어 사본으로 읽음: {tmp}")
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True)


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else latest_workbook()
    print("워크북:", path.name)
    wb = open_workbook(path)
    ws = wb["Exports_RawData"]
    rows = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        year, region, cen, ckr, hs, usd_k = r[1], r[2], r[3], r[4], r[5], r[6]
        if not isinstance(year, (int, float)) or not cen or usd_k is None:
            continue
        rows.append({"year": int(year), "region": str(region).strip(),
                     "country_en": str(cen).strip(),
                     "country_kr": str(ckr or "").strip(),
                     "hs": str(hs).strip(), "usd_k": float(usd_k)})
    df = pd.DataFrame(rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False, encoding="utf-8")
    yr = df.groupby("year")["usd_k"].sum() / 1e6
    print(f"customs_exports.csv: {len(df)}행, {df['year'].min()}~{df['year'].max()}"
          f", 국가 {df['country_en'].nunique()}개, HS {sorted(df['hs'].unique())}")
    print("연도 합계 (US$bn):")
    for y, v in yr.items():
        print(f"  {y}: {v:,.2f}")


if __name__ == "__main__":
    main()
